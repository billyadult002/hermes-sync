#!/usr/bin/env python3
from __future__ import annotations

import json
import re
import sys
import time
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt


ALLOWED_FORMATS = {"txt", "md", "json", "docx", "xlsx", "pptx", "rtf", "csv", "pages", "numbers", "key"}
APPLE_COMPATIBLE_FORMATS = {
    "pages": "docx",
    "numbers": "xlsx",
    "key": "pptx",
}


def slugify(value: str) -> str:
    text = re.sub(r"[^A-Za-z0-9\u4e00-\u9fff._-]+", "-", str(value or "").strip())
    text = re.sub(r"-{2,}", "-", text).strip("-._")
    return text or "workspace-output"


def bullet_lines(text: str) -> list[str]:
    items = []
    for raw in str(text or "").splitlines():
        line = raw.strip()
        if line:
            items.append(line)
    return items or ["No content available."]


def write_text(path: Path, payload: dict, markdown: bool = False) -> None:
    heading = payload["title"]
    body = payload["assistant_text"].strip() or "No content available."
    prompt = payload["prompt"].strip()
    attachments = payload.get("attachments", [])
    skill = payload["skill_id"]
    workspace = payload["workspace"]
    created_at = payload["created_at"]
    if markdown:
        parts = [
            f"# {heading}",
            "",
            f"- Workspace: `{workspace}`",
            f"- Skill: `{skill}`",
            f"- Created At: `{created_at}`",
        ]
        if prompt:
            parts.extend(["", "## Request", "", prompt])
        parts.extend(["", "## Result", "", body])
        if attachments:
            parts.extend(["", "## Attachments", ""])
            for item in attachments:
                parts.append(f"- {item.get('name') or 'Attachment'}")
        path.write_text("\n".join(parts), encoding="utf-8")
        return

    parts = [
        heading,
        f"Workspace: {workspace}",
        f"Skill: {skill}",
        f"Created At: {created_at}",
    ]
    if prompt:
        parts.extend(["", "Request", prompt])
    parts.extend(["", "Result", body])
    if attachments:
        parts.extend(["", "Attachments"])
        parts.extend(f"- {item.get('name') or 'Attachment'}" for item in attachments)
    path.write_text("\n".join(parts), encoding="utf-8")


def write_json(path: Path, payload: dict) -> None:
    serializable = {
        "title": payload["title"],
        "workspace": payload["workspace"],
        "skill_id": payload["skill_id"],
        "created_at": payload["created_at"],
        "prompt": payload["prompt"],
        "assistant_text": payload["assistant_text"],
        "attachments": payload.get("attachments", []),
    }
    path.write_text(json.dumps(serializable, ensure_ascii=False, indent=2), encoding="utf-8")


def rtf_escape(text: str) -> str:
    escaped = str(text or "").replace("\\", "\\\\").replace("{", "\\{").replace("}", "\\}")
    return escaped.replace("\n", "\\par\n")


def write_rtf(path: Path, payload: dict) -> None:
    body = payload["assistant_text"].strip() or "No content available."
    prompt = payload["prompt"].strip()
    parts = [
        r"{\rtf1\ansi\deff0",
        r"{\fonttbl{\f0 Arial;}}",
        r"\fs28\b " + rtf_escape(payload["title"]) + r"\b0\par",
        r"\fs20 Workspace: " + rtf_escape(payload["workspace"]) + r"\par",
        r"Skill: " + rtf_escape(payload["skill_id"]) + r"\par",
        r"Created At: " + rtf_escape(payload["created_at"]) + r"\par\par",
    ]
    if prompt:
        parts.extend([r"\b Request\b0\par", rtf_escape(prompt), r"\par\par"])
    parts.extend([r"\b Result\b0\par", rtf_escape(body), r"\par"])
    if payload.get("attachments"):
        parts.extend([r"\par\b Attachments\b0\par"])
        for item in payload["attachments"]:
            parts.append(r"- " + rtf_escape(item.get("name") or "Attachment") + r"\par")
    parts.append("}")
    path.write_text("\n".join(parts), encoding="utf-8")


def write_csv(path: Path, payload: dict) -> None:
    import csv

    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["Field", "Value"])
        writer.writerow(["Title", payload["title"]])
        writer.writerow(["Workspace", payload["workspace"]])
        writer.writerow(["Skill", payload["skill_id"]])
        writer.writerow(["Created At", payload["created_at"]])
        writer.writerow([])
        writer.writerow(["Request", payload["prompt"].strip()])
        writer.writerow([])
        writer.writerow(["Result"])
        for line in bullet_lines(payload["assistant_text"]):
            writer.writerow([line])
        if payload.get("attachments"):
            writer.writerow([])
            writer.writerow(["Attachments"])
            for item in payload["attachments"]:
                writer.writerow([item.get("name") or "Attachment", item.get("size") or ""])


def write_docx(path: Path, payload: dict) -> None:
    doc = Document()
    doc.add_heading(payload["title"], level=1)
    meta = doc.add_table(rows=3, cols=2)
    meta.rows[0].cells[0].text = "Workspace"
    meta.rows[0].cells[1].text = payload["workspace"]
    meta.rows[1].cells[0].text = "Skill"
    meta.rows[1].cells[1].text = payload["skill_id"]
    meta.rows[2].cells[0].text = "Created At"
    meta.rows[2].cells[1].text = payload["created_at"]

    if payload["prompt"].strip():
        doc.add_heading("Request", level=2)
        doc.add_paragraph(payload["prompt"].strip())

    doc.add_heading("Result", level=2)
    for line in bullet_lines(payload["assistant_text"]):
        doc.add_paragraph(line)

    attachments = payload.get("attachments", [])
    if attachments:
        doc.add_heading("Attachments", level=2)
        for item in attachments:
            doc.add_paragraph(str(item.get("name") or "Attachment"), style="List Bullet")
    doc.save(path)


def write_xlsx(path: Path, payload: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Title"
    ws["B1"] = payload["title"]
    ws["A2"] = "Workspace"
    ws["B2"] = payload["workspace"]
    ws["A3"] = "Skill"
    ws["B3"] = payload["skill_id"]
    ws["A4"] = "Created At"
    ws["B4"] = payload["created_at"]
    ws["A6"] = "Request"
    ws["B6"] = payload["prompt"].strip()
    ws["A8"] = "Result"
    for idx, line in enumerate(bullet_lines(payload["assistant_text"]), start=9):
        ws.cell(row=idx, column=2).value = line
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 120

    att_ws = wb.create_sheet("Attachments")
    att_ws.append(["Name", "Size", "Preview"])
    for item in payload.get("attachments", []):
        att_ws.append([
            item.get("name") or "",
            item.get("size") or "",
            str(item.get("preview") or "")[:1000],
        ])
    att_ws.column_dimensions["A"].width = 40
    att_ws.column_dimensions["B"].width = 14
    att_ws.column_dimensions["C"].width = 100
    wb.save(path)


def write_pptx(path: Path, payload: dict) -> None:
    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = payload["title"]
    title_slide.placeholders[1].text = f"{payload['workspace']} · {payload['skill_id']} · {payload['created_at']}"

    body_slide = prs.slides.add_slide(prs.slide_layouts[1])
    body_slide.shapes.title.text = "Workspace Output"
    tf = body_slide.placeholders[1].text_frame
    tf.clear()
    prompt = payload["prompt"].strip()
    if prompt:
        p = tf.paragraphs[0]
        p.text = f"Request: {prompt[:200]}"
        p.font.size = Pt(18)
    for line in bullet_lines(payload["assistant_text"])[:8]:
        para = tf.add_paragraph()
        para.text = line[:220]
        para.level = 0
        para.font.size = Pt(20)

    if payload.get("attachments"):
        att_slide = prs.slides.add_slide(prs.slide_layouts[5])
        att_slide.shapes.title.text = "Attachments"
        box = att_slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.5), Inches(4.5))
        att_tf = box.text_frame
        for item in payload["attachments"]:
            para = att_tf.add_paragraph()
            para.text = str(item.get("name") or "Attachment")
            para.level = 0
            para.font.size = Pt(20)
    prs.save(path)


def main() -> int:
    payload = json.loads(sys.stdin.read() or "{}")
    fmt = str(payload.get("format") or "").lower().strip()
    if fmt not in ALLOWED_FORMATS:
        raise SystemExit(json.dumps({"ok": False, "error": f"unsupported_format:{fmt}"}))
    requested_fmt = fmt
    actual_fmt = APPLE_COMPATIBLE_FORMATS.get(fmt, fmt)

    output_dir = Path(payload["output_dir"]).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    title = str(payload.get("title") or "Workspace Output").strip()
    created_at = payload.get("created_at") or time.strftime("%Y-%m-%d %H:%M:%S")
    normalized = {
        "title": title,
        "workspace": str(payload.get("workspace") or "workspace"),
        "skill_id": str(payload.get("skill_id") or "general"),
        "prompt": str(payload.get("prompt") or ""),
        "assistant_text": str(payload.get("assistant_text") or ""),
        "attachments": payload.get("attachments") or [],
        "created_at": created_at,
    }
    filename = f"{slugify(title)}-{time.strftime('%Y%m%d-%H%M%S')}.{actual_fmt}"
    path = output_dir / filename

    if actual_fmt == "txt":
        write_text(path, normalized, markdown=False)
    elif actual_fmt == "md":
        write_text(path, normalized, markdown=True)
    elif actual_fmt == "json":
        write_json(path, normalized)
    elif actual_fmt == "docx":
        write_docx(path, normalized)
    elif actual_fmt == "xlsx":
        write_xlsx(path, normalized)
    elif actual_fmt == "pptx":
        write_pptx(path, normalized)
    elif actual_fmt == "rtf":
        write_rtf(path, normalized)
    elif actual_fmt == "csv":
        write_csv(path, normalized)

    print(json.dumps({
        "ok": True,
        "path": str(path),
        "filename": path.name,
        "format": requested_fmt,
        "actual_format": actual_fmt,
        "apple_compatible": requested_fmt in APPLE_COMPATIBLE_FORMATS,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
