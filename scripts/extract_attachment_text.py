#!/usr/bin/env python3
from __future__ import annotations

import base64
import io
import json
import re
import sys
import zipfile
from xml.etree import ElementTree as ET

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


MAX_TEXT = 8000


def compact_text(text: str, limit: int = MAX_TEXT) -> str:
    if not text:
        return ""
    lines = []
    for raw in str(text).replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = " ".join(raw.split()).strip()
        if line:
            lines.append(line)
    out = []
    total = 0
    for line in lines:
        addition = len(line) + 1
        if total + addition > limit:
            break
        out.append(line)
        total += addition
    return "\n".join(out)


def likely_metric_line(line: str) -> bool:
    score = 0
    if any(token in line.lower() for token in ("revenue", "ebitda", "profit", "margin", "cash", "debt", "asset", "liability", "收入", "利润", "毛利", "现金", "债务", "资产", "负债", "%")):
        score += 1
    if sum(ch.isdigit() for ch in line) >= 3:
        score += 1
    if any(token in line for token in ("$", "¥", "%", "亿元", "万元", "million", "billion", "x", "同比", "YoY", "QoQ")):
        score += 1
    return score >= 2


def summarize_lines(lines: list[str], limit: int = 10) -> list[str]:
    picked = []
    seen = set()
    for line in lines:
        normalized = line.strip()
        if not normalized or normalized in seen:
            continue
        if likely_metric_line(normalized):
            picked.append(normalized)
            seen.add(normalized)
        if len(picked) >= limit:
            break
    return picked


def pick_signal_lines(lines: list[str], keywords: tuple[str, ...], limit: int = 6) -> list[str]:
    picked = []
    seen = set()
    lowered_keywords = tuple(keyword.lower() for keyword in keywords)
    for line in lines:
        normalized = line.strip()
        if not normalized or normalized in seen:
            continue
        lowered = normalized.lower()
        if any(keyword in lowered for keyword in lowered_keywords):
            picked.append(normalized)
            seen.add(normalized)
        if len(picked) >= limit:
            break
    return picked


def finance_blocks(preview: str) -> list[str]:
    lines = preview.splitlines()
    metrics = summarize_lines(lines, limit=10)
    positives = pick_signal_lines(lines, ("increase", "grew", "improved", "up", "增长", "提升", "改善", "增加", "强劲"))
    risks = pick_signal_lines(lines, ("risk", "decline", "decrease", "down", "debt", "covenant", "pressure", "uncertain", "风险", "下降", "债务", "承压", "不确定"))
    blocks: list[str] = []
    if metrics:
        blocks.extend(["[Possible Core Metrics]"] + metrics + [""])
    if positives:
        blocks.extend(["[Possible Positive Signals]"] + positives + [""])
    if risks:
        blocks.extend(["[Possible Risk Signals]"] + risks + [""])
    if preview:
        blocks.extend(["[Extract Preview]"] + lines[:60])
    return blocks


def financial_report_summary_blocks(preview: str) -> list[str]:
    lines = preview.splitlines()
    metrics = summarize_lines(lines, limit=12)
    positives = pick_signal_lines(lines, ("increase", "grew", "improved", "up", "beat", "增长", "提升", "改善", "增加", "超预期", "强劲"))
    risks = pick_signal_lines(lines, ("risk", "decline", "decrease", "down", "debt", "covenant", "pressure", "uncertain", "loss", "风险", "下降", "债务", "承压", "不确定", "亏损"))
    outlook = pick_signal_lines(lines, ("guidance", "outlook", "forecast", "expect", "预计", "展望", "指引", "目标", "管理层"))
    blocks: list[str] = []
    blocks.append("【附件提取底稿】")
    if metrics:
        blocks.extend(["", "【核心财务数据】"] + [f"- {line}" for line in metrics[:8]])
    if positives:
        blocks.extend(["", "【正面信号】"] + [f"- {line}" for line in positives[:6]])
    if risks:
        blocks.extend(["", "【风险提示】"] + [f"- {line}" for line in risks[:6]])
    if outlook:
        blocks.extend(["", "【管理层指引与展望】"] + [f"- {line}" for line in outlook[:6]])
    blocks.extend(["", "【待核实摘录】"] + [f"- {line}" for line in lines[:20]])
    return blocks


def extract_pdf(data: bytes, skill_id: str = "") -> str:
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages[:10]:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
        if sum(len(item) for item in parts) >= MAX_TEXT:
            break
    preview = compact_text("\n".join(parts))
    if skill_id == "financial-report-summary":
        blocks = financial_report_summary_blocks(preview)
    elif skill_id in {"financial-analysis", "am_reporting_engine"}:
        blocks = finance_blocks(preview)
    else:
        lines = preview.splitlines()
        metrics = summarize_lines(lines)
        blocks = []
        if metrics:
            blocks.extend(["[Detected Key Lines]"] + metrics + [""])
        if preview:
            blocks.extend(["[Extract Preview]"] + lines[:60])
    return "\n".join(blocks)[:MAX_TEXT]


def extract_docx(data: bytes, skill_id: str = "") -> str:
    doc = Document(io.BytesIO(data))
    parts = [para.text for para in doc.paragraphs if para.text.strip()]
    preview = compact_text("\n".join(parts))
    if skill_id == "financial-report-summary":
        blocks = financial_report_summary_blocks(preview)
    elif skill_id in {"financial-analysis", "am_reporting_engine", "client-follow-up-weekly-update"}:
        blocks = finance_blocks(preview)
    else:
        lines = preview.splitlines()
        metrics = summarize_lines(lines)
        blocks = []
        if metrics:
            blocks.extend(["[Detected Key Lines]"] + metrics + [""])
        if preview:
            blocks.extend(["[Extract Preview]"] + lines[:60])
    return "\n".join(blocks)[:MAX_TEXT]


def extract_xlsx(data: bytes, skill_id: str = "") -> str:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets[:3]:
        lines.append(f"[Sheet] {sheet.title}")
        for row in sheet.iter_rows(max_row=30, values_only=True):
            values = [str(cell).strip() for cell in row if cell not in (None, "")]
            if values:
                lines.append(" | ".join(values))
            if sum(len(item) for item in lines) >= MAX_TEXT:
                break
        if sum(len(item) for item in lines) >= MAX_TEXT:
            break
    metrics = summarize_lines(lines, limit=12)
    blocks = []
    if skill_id == "financial-report-summary":
        if metrics:
            blocks.extend(["【附件提取底稿】", "", "【核心财务数据】"] + [f"- {line}" for line in metrics[:10]])
        risk_rows = pick_signal_lines(lines, ("debt", "covenant", "risk", "liability", "下降", "风险", "债务", "负债"), limit=8)
        if risk_rows:
            blocks.extend(["", "【风险提示】"] + [f"- {line}" for line in risk_rows[:8]])
        blocks.extend(["", "【待核实摘录】"] + [f"- {line}" for line in lines[:30]])
    elif skill_id in {"financial-analysis", "pdf-excel-analysis", "am_reporting_engine"}:
        if metrics:
            blocks.extend(["[Possible Core Rows]"] + metrics + [""])
        risk_rows = pick_signal_lines(lines, ("debt", "covenant", "risk", "liability", "下降", "风险", "债务", "负债"), limit=8)
        if risk_rows:
            blocks.extend(["[Possible Risk Rows]"] + risk_rows + [""])
    else:
        if metrics:
            blocks.extend(["[Detected Key Rows]"] + metrics + [""])
    blocks.extend(["[Workbook Preview]"] + lines[:80])
    return "\n".join(blocks)[:MAX_TEXT]


def extract_pptx(data: bytes, skill_id: str = "") -> str:
    prs = Presentation(io.BytesIO(data))
    lines = []
    for idx, slide in enumerate(prs.slides, start=1):
        if idx > 12:
            break
        lines.append(f"[Slide] {idx}")
        slide_texts = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and str(text).strip():
                slide_texts.append(str(text).strip())
        lines.extend(slide_texts[:12])
        if sum(len(item) for item in lines) >= MAX_TEXT:
            break
    preview = compact_text("\n".join(lines))
    key_lines = summarize_lines(preview.splitlines(), limit=10)
    blocks = []
    if key_lines:
        blocks.extend(["[Slide Highlights]"] + key_lines + [""])
    if skill_id == "financial-report-summary":
        blocks = financial_report_summary_blocks(preview)
    elif skill_id in {"financial-analysis", "client-follow-up-weekly-update", "mckinsey-thinking"}:
        blocks = finance_blocks(preview)
    else:
        blocks.extend(["[Deck Preview]"] + preview.splitlines()[:80])
    return "\n".join(blocks)[:MAX_TEXT]


def extract_iwork(data: bytes, filename: str, skill_id: str = "") -> str:
    lower = filename.lower()
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
            for preview_name in ("QuickLook/Preview.pdf", "preview.pdf", "Preview.pdf"):
                if preview_name in names:
                    return extract_pdf(zf.read(preview_name), skill_id)

            text_parts = []
            xml_candidates = [name for name in names if name.lower().endswith(".xml")][:8]
            for name in xml_candidates:
                try:
                    root = ET.fromstring(zf.read(name))
                    texts = [node.strip() for node in root.itertext() if node and node.strip()]
                    if texts:
                        text_parts.extend(texts[:200])
                except Exception:
                    continue

            txt_candidates = [name for name in names if name.lower().endswith((".txt", ".csv", ".tsv"))][:8]
            for name in txt_candidates:
                try:
                    text_parts.append(zf.read(name).decode("utf-8", errors="ignore"))
                except Exception:
                    continue

            preview = compact_text("\n".join(text_parts))
            if preview:
                header = "[iWork Preview]" if lower.endswith((".pages", ".key")) else "[iWork Workbook Preview]"
                if skill_id == "financial-report-summary":
                    return "\n".join(financial_report_summary_blocks(preview))[:MAX_TEXT]
                if skill_id in {"financial-analysis", "pdf-excel-analysis", "client-follow-up-weekly-update", "mckinsey-thinking"}:
                    return "\n".join(finance_blocks(preview))[:MAX_TEXT]
                return "\n".join([header] + preview.splitlines()[:80])[:MAX_TEXT]
    except Exception:
        return ""
    return ""


def extract_rtf(data: bytes, skill_id: str = "") -> str:
    raw = data.decode("utf-8", errors="ignore")
    text = raw
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
    text = re.sub(r"[{}]", " ", text)
    preview = compact_text(text)
    if skill_id == "financial-report-summary":
        return "\n".join(financial_report_summary_blocks(preview))[:MAX_TEXT]
    if skill_id in {"financial-analysis", "client-follow-up-weekly-update", "document-review"}:
        return "\n".join(finance_blocks(preview))[:MAX_TEXT]
    return "\n".join(["[RTF Preview]"] + preview.splitlines()[:80])[:MAX_TEXT]


def extract_text(data: bytes, filename: str, mime: str, skill_id: str = "") -> str:
    lower = filename.lower()
    if mime.startswith("text/") or lower.endswith((".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml")):
        preview = compact_text(data.decode("utf-8", errors="ignore"))
        if skill_id == "financial-report-summary":
            return "\n".join(financial_report_summary_blocks(preview))[:MAX_TEXT]
        if skill_id in {"financial-analysis", "client-follow-up-weekly-update"}:
            return "\n".join(finance_blocks(preview))[:MAX_TEXT]
        return "\n".join(["[Text Preview]"] + preview.splitlines()[:80])[:MAX_TEXT]
    if lower.endswith(".pdf") or mime == "application/pdf":
        return extract_pdf(data, skill_id)
    if lower.endswith((".docx", ".docm")) or mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extract_docx(data, skill_id)
    if lower.endswith((".xlsx", ".xlsm")) or mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return extract_xlsx(data, skill_id)
    if lower.endswith((".pptx", ".pptm")) or mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        return extract_pptx(data, skill_id)
    if lower.endswith((".pages", ".numbers", ".key")):
        return extract_iwork(data, filename, skill_id)
    if lower.endswith(".rtf") or mime in {"application/rtf", "text/rtf"}:
        return extract_rtf(data, skill_id)
    return ""


def main() -> int:
    payload = json.loads(sys.stdin.read() or "{}")
    filename = str(payload.get("filename") or "attachment")
    mime = str(payload.get("mime") or "")
    skill_id = str(payload.get("skill_id") or "")
    blob = str(payload.get("base64") or "")
    raw = base64.b64decode(blob.encode("utf-8")) if blob else b""
    preview = extract_text(raw, filename, mime, skill_id)
    print(json.dumps({"ok": True, "preview": preview}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
